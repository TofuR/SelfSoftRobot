#!/usr/bin/env python3
"""Export the Chinese manuscript to editable Word with numbered native math.

Requires pandoc and python-docx. The Markdown source is never modified.
"""
from __future__ import annotations

import argparse
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent


def portable_math(source: str) -> str:
    """Preserve math meaning while avoiding older Office renderer ambiguities."""
    def convert(match: re.Match) -> str:
        display = match.group(1) is not None
        value = match.group(1) if display else match.group(2)
        value = re.sub(
            r"\\begin\{aligned\}(.*?)\\end\{aligned\}",
            lambda m: r"\begin{gathered}" + m[1].replace("&", "") + r"\end{gathered}",
            value, flags=re.S,
        )
        # Pandoc drops tag labels in native Word math; make each label explicit.
        value = re.sub(r"\\tag\{(\d+)\}", lambda m: r"\qquad (" + m[1] + ")", value)
        # Older LibreOffice interprets the function named vec as a vector accent.
        value = re.sub(r"\\operatorname\{([^}]+)\}", r"\\mathrm{\1}", value)
        delimiter = "$$" if display else "$"
        return delimiter + value + delimiter

    return re.sub(r"\$\$(.*?)\$\$|\$(?!\$)([^\n$]+)\$", convert, source, flags=re.S)


def style_document(path: Path) -> None:
    doc = Document(path)
    for section in doc.sections:
        section.page_height, section.page_width = Cm(29.7), Cm(21)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2)
        section.header_distance = section.footer_distance = Cm(.8)
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            style.font.name, style.font.size = "Times New Roman", Pt(11)
            fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                style.element.get_or_add_rPr().append(fonts)
            fonts.set(qn("w:eastAsia"), "Noto Serif CJK SC")
    for name in ("Normal", "Body Text", "First Paragraph"):
        if name in doc.styles:
            fmt = doc.styles[name].paragraph_format
            fmt.line_spacing, fmt.space_after = 1.15, Pt(6)
    for name, size in (("Title", 19), ("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.size, style.font.bold = Pt(size), True
        style.font.color.rgb = RGBColor.from_string("153C40")
        style.paragraph_format.keep_with_next = True
        style.element.get_or_add_rPr().find(qn("w:rFonts")).set(qn("w:eastAsia"), "Noto Sans CJK SC")
    if "Draft Table" not in doc.styles:
        doc.styles.add_style("Draft Table", WD_STYLE_TYPE.TABLE)
    for table in doc.tables:
        table.style, table.autofit = "Draft Table", False
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement("w:" + edge)
            for key, value in (("val", "single"), ("sz", "4"), ("color", "BCC9CB")):
                element.set(qn("w:" + key), value)
            borders.append(element)
        table._tbl.tblPr.append(borders)
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for cell in row.cells:
                cell.width = Cm(17 / len(row.cells))
                if row_index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "EAF0F0")
                    cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    fmt = paragraph.paragraph_format
                    fmt.space_after = fmt.space_before = Pt(3)
                    fmt.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5 if len(row.cells) > 5 else 9)
                        if row_index == 0:
                            run.font.bold = True
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = "ICRA 2027 · 中文工作初稿 · 实验结果待补"
        for run in header.runs:
            run.font.size, run.font.color.rgb = Pt(8), RGBColor(100, 100, 100)
        footer = section.footer.paragraphs[0]
        footer.alignment = 2
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    doc.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=BASE / "manuscript_zh.docx")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    if args.out.exists() and not args.overwrite:
        parser.error("output exists; save unique Word comments before using --overwrite")
    source = portable_math((BASE / "manuscript_zh.md").read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory(prefix="icra2027-word-") as folder:
        markdown = Path(folder) / "manuscript.md"
        output = Path(folder) / "manuscript.docx"
        markdown.write_text(source, encoding="utf-8")
        subprocess.run([
            "pandoc", str(markdown), "--from=markdown+tex_math_dollars",
            "--to=docx", "--standalone", "--output=" + str(output),
        ], check=True)
        style_document(output)
        args.out.parent.mkdir(parents=True, exist_ok=True)
        with args.out.open("wb" if args.overwrite else "xb") as destination:
            destination.write(output.read_bytes())
    print(args.out)


if __name__ == "__main__":
    main()
